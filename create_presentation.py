from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()

# Slide 1: Title Slide
title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Taco Tequilas: Savoring the Code"
subtitle.text = "A Vibe Coding Journey to Authentic Flavor\n\nPresenter: Freddie Porter"
# Notes
slide.notes_slide.notes_text_frame.text = "Welcome, everyone. Today I'm going to walk you through how we brought the authentic Mexican flavor of Taco Tequilas to the web using modern, AI-assisted development techniques."

# Slide 2: Overall Problem
bullet_slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "The Challenge: Capturing the Fiesta Online"
tf = body_shape.text_frame
tf.text = "A local favorite needing a digital facelift"
p = tf.add_paragraph()
p.text = "Customers struggled to find accurate menus and locations"
p = tf.add_paragraph()
p.text = "The need for speed: Building rapidly without sacrificing quality"
# Notes
slide.notes_slide.notes_text_frame.text = "Imagine craving authentic tacos and a perfectly mixed margarita, but you can't figure out if the local spot is open or where they are located. That was the problem Taco Tequilas faced. They had amazing food, but a digital presence that left customers hungry for information. We needed to craft a modern, responsive website that looked as good as the food tastes—and we needed to do it incredibly fast. Our goal was to create a seamless way for users to explore the menu and find the nearest location, ensuring the digital vibe matched the in-person fiesta."

# Slide 3: Software Stack
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "The 'Vibe Coding' Stack"
tf = body_shape.text_frame
tf.text = "Editor: VS Code"
p = tf.add_paragraph()
p.text = "Generative AI: GitHub Copilot"
p = tf.add_paragraph()
p.text = "Automated Testing: Antigravity (Agentic AI)"
p = tf.add_paragraph()
p.text = "Deployment: Netlify & GitHub"
# Notes
slide.notes_slide.notes_text_frame.text = "To achieve this rapid, high-quality build, we embraced what we call the 'Vibe Coding' stack. I used VS Code with Copilot for rapid HTML, CSS, and JS generation, letting the AI handle the boilerplate so I could focus on the design and user experience. But generation is only half the battle. To ensure reliability, I used Antigravity for automated end-to-end testing."

# Slide 4: Critical Design Tasks
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "Structuring the User Experience"
tf = body_shape.text_frame
tf.text = "Task 1: Central Hub Home Page (Nav, Hero, Visit, Footer)"
p = tf.add_paragraph()
p.text = "Task 2: Interactive Maps (Google Maps API for locations)"
p = tf.add_paragraph()
p.text = "Task 3: Persistent Facebook Footer (Live updates)"
p = tf.add_paragraph()
p.text = "Task 4: 'About' Page (Origin, values, milestones)"
p = tf.add_paragraph()
p.text = "Task 5: Transparent Menu (Prices, descriptions, imagery)"
# Notes
slide.notes_slide.notes_text_frame.text = "Our design focused on five critical user tasks. First, the Home page serves as a simple, low-effort central hub. Second, we integrated the Google Maps API for interactive, zoomable location previews. Third, we placed a persistent Facebook link in the footer across all pages for live updates without disrupting navigation. Fourth, we built a Story section to establish credibility by sharing the restaurant's origin and values. Finally, Task 5 centered on a transparent menu experience, clearly presenting dishes with pricing and descriptions."

# Slide 5: Testing
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "Trust, but Verify: Catching AI Hallucinations"
tf = body_shape.text_frame
tf.text = "Process: Automated UI navigation and link validation using Antigravity"
p = tf.add_paragraph()
p.text = "The 'Hallucination': Copilot generated relative URLs for external social links."
p = tf.add_paragraph()
p.text = "The Fix: Correcting protocols and adding security attributes."
# Notes
slide.notes_slide.notes_text_frame.text = "Our testing process involved having Antigravity autonomously click through the live site, verifying every menu link and footer icon. This is where we learned a valuable lesson about AI generation. Copilot wrote this code for our social footer: <a href=\"facebook.com/TacoTequilaTH\">. It looked correct at a glance, but Antigravity failed the test because clicking it caused the browser to look for a local page, resulting in a 404 error. Here is how I fixed it: I updated the links to include the full https:// protocol and added target=\"_blank\" rel=\"noopener noreferrer\" to safely open the links in a new tab. It proved that while Copilot is fast, automated testing is essential to catch subtle hallucinations."

# Slide 6: Prototype Walkthrough
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "The Live Experience: Taco Tequilas App"
tf = body_shape.text_frame
tf.text = "Navigating the Home Page Hub (Task 1 & 3)"
p = tf.add_paragraph()
p.text = "Exploring the Menu (Task 5)"
p = tf.add_paragraph()
p.text = "Locating a Restaurant (Task 2)"
p = tf.add_paragraph()
p.text = "Discovering the Story (Task 4)"
# Notes
slide.notes_slide.notes_text_frame.text = "Let's walk through the live application. First, notice the Home page acts as our central hub, giving you quick access to everything, with the Facebook footer persisting at the bottom for live updates. \nIf we click 'Menu', you can see our transparent pricing and item descriptions, fulfilling Task 5. \nNext, a customer wants to know where to go. We click 'Visit' in the navigation. Here, we see interactive Google Maps. Users can pan and zoom directly on the page, fulfilling Task 2. \nFinally, we can visit the 'Story' page to learn about the origins and milestones of Taco Tequilas, completing Task 4."

# Slide 7: Summary
slide = prs.slides.add_slide(bullet_slide_layout)
shapes = slide.shapes
title_shape = shapes.title
body_shape = shapes.placeholders[1]
title_shape.text = "The Power of AI Collaboration"
tf = body_shape.text_frame
tf.text = "Vibe Coding accelerates the boilerplate"
p = tf.add_paragraph()
p.text = "Automated agents are crucial for quality assurance"
p = tf.add_paragraph()
p.text = "The surprise: How natural it feels to 'pair program' with an AI"
# Notes
slide.notes_slide.notes_text_frame.text = "To summarize, designing this application taught me that the future of web development is deeply collaborative with AI. I thoroughly enjoyed using these tools—they removed the tedium and let me focus on the creative aspects of the site. What surprised me the most was how effectively Copilot and Antigravity complemented each other. Copilot acts as the enthusiastic builder, while Antigravity acts as the meticulous inspector. Together, they made it possible to deliver a premium, fully functional prototype covering all 5 of our core tasks in record time. Thank you."

prs.save('/Users/freddieporter/repos/TacoTequilas/TacoTequilas_Presentation.pptx')
